"""Memory-only CSV/XLSX import and evaluation without exporting query text."""
import hashlib
import math
from collections import defaultdict
import csv
import io
import json
import zipfile
from itertools import islice
from pathlib import Path
from .registry import REGISTRY
from .corpus import source_manifest, PROVENANCE, STATUSES, SCOPES


def load_demo_suite():
    path = Path(__file__).resolve().parents[1] / "data" / "query_suite.json"
    cases = json.loads(path.read_text())
    for c in cases:
        c.update(provenance="rewritten_fixture", annotation_status="proposed", independent_review=False,
                 label_semantics="required", group=f"source_group_{c['source_row']}", source_rows=[c['source_row']],
                 quarantined=False, quarantine_reasons=[], annotator_role="registry_author")
    return cases


def import_suite(payload: bytes, filename: str) -> list[dict]:
    if len(payload) > 5 * 1024 * 1024:
        raise ValueError("Upload a file smaller than 5 MB.")
    if filename.lower().endswith(".csv"):
        rows = list(islice(csv.reader(io.StringIO(payload.decode("utf-8-sig"))), 2002))
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            if sum(i.file_size for i in z.infolist()) > 25 * 1024 * 1024 or len(z.infolist()) > 2000:
                raise ValueError("Expanded workbook exceeds the POC size limit.")
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        sheet = workbook["Visible Rows"] if "Visible Rows" in workbook.sheetnames else workbook.active
        # Bound malicious/inflated worksheet dimensions as well as ZIP size.
        rows = list(islice(sheet.iter_rows(values_only=True, max_col=30), 2002))
        workbook.close()
    else:
        raise ValueError("Supported formats are CSV and XLSX.")
    if not rows:
        raise ValueError("The file has no rows.")
    if len(rows) > 2001:
        raise ValueError("The POC supports at most 2,000 physical data rows including blanks.")
    headers = [str(v or "").strip().lower() for v in rows[0]]
    qname = next((x for x in ["query", "advisor query example"] if x in headers), None)
    if qname is None:
        raise ValueError("The first row must include 'query' or 'Advisor Query Example'.")
    qi = headers.index(qname)
    ei = headers.index("expected_intents") if "expected_intents" in headers else None
    cases = []
    manifest = source_manifest()
    known_source = hashlib.sha256(payload).hexdigest() == manifest["source_sha256"]
    for physical_row, row in enumerate(rows[1:],2):
        if qi >= len(row) or row[qi] is None or not str(row[qi]).strip():
            continue
        query = str(row[qi]).strip()
        if len(query) > 16000:
            raise ValueError("A query exceeds the 16,000-character limit.")
        expected = str(row[ei] or "").replace(";", ",").split(",") if ei is not None and ei < len(row) else []
        expected = [x.strip() for x in expected if x.strip()]
        if any(x not in REGISTRY for x in expected):
            raise ValueError("An expected_intents value is not in the intent registry.")
        def field(name, default=""):
            return str(row[headers.index(name)] or default).strip() if name in headers and headers.index(name)<len(row) else default
        provenance = field("provenance", "unknown")
        status = field("annotation_status", "proposed" if expected else "unlabeled")
        semantics = field("label_semantics", "required")
        scope = field("expected_scope")
        expected_decision = field("expected_decision")
        if expected_decision not in {'','ROUTE_PREVIEW','CLARIFY','REVIEW'}:
            raise ValueError("Invalid expected decision.")
        if provenance not in PROVENANCE or status not in STATUSES or semantics not in {"required","exhaustive"} or scope not in SCOPES:
            raise ValueError("Invalid annotation metadata.")
        note_present = bool(field("transcription notes"))
        thread = next((f"source_thread_{g[0]}" for g in manifest["suggested_threads"] if physical_row in g), None) if known_source else None
        group = thread or f"source_group_{physical_row}"
        cases.append({"id": f"uploaded_row_{physical_row}", "source_row": physical_row, "source_rows":[physical_row],
                      "query": query, "expected_intents": expected, "provenance":provenance, "annotation_status":status,
                      "independent_review": field("independent_review").lower() == "true", "label_semantics":semantics,
                      "expected_scope":scope, "group":group, "suggested_thread":thread,
                      "expected_decision":expected_decision,
                      "quarantined": note_present or (known_source and physical_row in manifest["notes_rows"]),
                      "quarantine_reasons": ["incomplete_source_transcription"] if note_present or (known_source and physical_row in manifest["notes_rows"]) else [],
                      "origin":"original_source_workbook" if known_source else "uploaded_unverified_provenance"})
    if len(cases) > 500:
        raise ValueError("The POC supports at most 500 nonempty queries per batch.")
    if not cases:
        raise ValueError("No nonempty queries were found.")
    return cases


def _wilson(successes, n):
    if not n:
        return None
    z = 1.96
    p = successes/n
    center = (p+z*z/(2*n))/(1+z*z/n)
    radius = z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/(1+z*z/n)
    return [round(max(0,center-radius),3),round(min(1,center+radius),3)]


def evaluate(cases, classifier, threshold=0.23, margin=0.035, sensitive_terms=(), compare_baseline=True):
    output, buckets = [], defaultdict(list)
    for n,case in enumerate(cases,1):
        if case.get("quarantined",False):
            output.append({"case":f"case_{n:03d}", "mode":"excluded", "decision":"QUARANTINED",
                           "predicted":"", "expected":"", "missing_expected":"", "additional_labels":"",
                           "metric_eligible":False, "provenance":case.get("provenance","unknown")})
            continue
        for mode in (["hybrid","rules_only"] if compare_baseline else ["hybrid"]):
            r = classifier.classify(case["query"], threshold=threshold, margin=margin, sensitive_terms=sensitive_terms, mode=mode)
            expected, predicted = set(case.get("expected_intents", [])), set(r["intents"])
            provenance = case.get("provenance","unknown")
            # Metadata is asserted by the operator; the POC cannot certify annotator independence.
            eligible = (case.get("annotation_status") == "adjudicated" and case.get("independent_review") is True
                        and provenance != "rewritten_fixture" and case.get("origin") != "manually_rewritten_from_source_scenario")
            exhaustive = eligible and case.get("label_semantics") == "exhaustive"
            length = len(case["query"])
            band = "under_120" if length < 120 else ("over_400" if length > 400 else "120_to_400")
            row = {"case":f"case_{n:03d}", "mode":mode, "decision":r["decision"],
                   "predicted":", ".join(r["intents"]), "expected":", ".join(sorted(expected)),
                   "missing_expected":", ".join(sorted(expected-predicted)), "additional_labels":", ".join(sorted(predicted-expected)) if expected else "",
                   "required_labels_found":expected <= predicted if expected or exhaustive else None,
                   "exact_label_set":expected == predicted if exhaustive else None,
                   "metric_eligible":eligible, "exhaustive_labels":exhaustive,
                   "provenance":provenance, "length_band":band, "characters":length,
                   "annotation_status":case.get("annotation_status","unlabeled"),
                   "composition":"compound" if len(expected)>1 else ("single" if expected else "unknown"),
                   "intent_count":len(predicted), "expected_count":len(expected), "found_count":len(expected&predicted),
                   "scope_correct":r["scope"]==case.get("expected_scope") if eligible and case.get("expected_scope") else None,
                   "decision_correct":r["decision"]==case.get("expected_decision") if eligible and case.get("expected_decision") else None,
                   "top_similarity":r["candidates"][0]["similarity"], "latency_ms":r["latency_ms"]}
            output.append(row)
            buckets[(mode,provenance,band,row["composition"])].append(row)
    strata=[]
    for (mode,provenance,band,composition),rows in sorted(buckets.items()):
        labeled=[r for r in rows if r['metric_eligible']]
        exhaustive=[r for r in labeled if r['exhaustive_labels']]
        denom=sum(r['expected_count'] for r in labeled)
        exact=sum(r['exact_label_set'] for r in exhaustive)
        scoped=[r for r in labeled if r['scope_correct'] is not None]
        decisions=[r for r in labeled if r['decision_correct'] is not None]
        strata.append({"mode":mode,"provenance":provenance,"length_band":band,"composition":composition,
                       "cases":len(rows),"reviewed_cases":len(labeled),"expected_label_assignments":denom,
                       "required_label_recall":round(sum(r['found_count'] for r in labeled)/denom,3) if denom else None,
                       "exact_set_rate":round(exact/len(exhaustive),3) if exhaustive else None,
                       "exact_set_wilson_95":_wilson(exact,len(exhaustive)),
                       "scope_accuracy":round(sum(r['scope_correct'] for r in scoped)/len(scoped),3) if scoped else None,
                       "decision_accuracy":round(sum(r['decision_correct'] for r in decisions)/len(decisions),3) if decisions else None,
                       "clarifications":sum(r['decision']=='CLARIFY' for r in rows)})
    main=[r for r in output if r['mode']=='hybrid']
    return {"rows":output,"strata":strata,"summary":{
        "cases":len(cases),"evaluated_cases":len(main),"quarantined_cases":sum(r['mode']=='excluded' for r in output),
        "reviewed_cases":sum(r['metric_eligible'] for r in main),"clarify_count":sum(r['decision']=='CLARIFY' for r in main),
        "review_count":sum(r['decision']=='REVIEW' for r in main),"outbound_llm_requests":0,
        "registry_digest":classifier.registry_digest,"settings":{"threshold":threshold,"margin":margin,"status":"unvalidated_placeholders"},
        "label_caveat":"No aggregate accuracy claim. Rewritten fixtures are regression diagnostics only. Metrics require operator-asserted independent, adjudicated labels; provenance unknown stays unknown. The screenshot corpus has no known production-frequency relationship."}}
