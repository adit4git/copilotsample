"""Turn an outside statement (PDF, scanned PDF, image, Excel, CSV) into structured holdings.

Pipeline
  1. read     - text PDF: pdftotext -layout; scanned PDF / image: page image -> tesseract (word boxes + confidence);
                Excel / CSV: cells.
  2. lines    - every source becomes the same thing: lines split into column segments, each with a confidence.
  3. parse    - metadata (institution, owner, account, as-of, stated total), the holdings table (header-driven column
                mapping), and any asset-allocation section.
  4. check    - quantity x price vs market value, sum of rows vs stated total, section subtotals, low-confidence and
                OCR-corrected cells.

Long statements (25+ pages of holdings, activity, charts and disclosures)
  - every page is classed first: summary, chart, holdings, activity, disclosure or other. Only summary and holdings
    pages are parsed for holdings; activity and fee tables never become holdings; disclosures are skipped.
  - a holdings table carries over page breaks, with or without a repeated column header; section subtotals are checked
    against their rows and do not end the table; a description that wraps onto a second line is kept with its row.
  - scanned PDFs are read in two passes: a fast low-resolution pass to class the pages, then full-resolution OCR only
    on the summary and holdings pages, in parallel.
  - allocations shown only as charts are measured (extract/charts.py).

Every number in the output comes from the statement. Nothing is inferred except a value computed as quantity x price
when the value column is missing, and that is flagged.
"""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

SUPPORTED = {'.pdf': 'pdf', '.png': 'image', '.jpg': 'image', '.jpeg': 'image', '.webp': 'image', '.tif': 'image', '.tiff': 'image',
             '.xlsx': 'xlsx', '.xlsm': 'xlsx', '.csv': 'csv', '.txt': 'text'}
MAX_FILES = 5
MAX_BYTES = 15 * 1024 * 1024

# ---------------------------------------------------------------------------------------------------------------- columns
FIELD_SYNONYMS = {
    'sym': ['symbol', 'ticker', 'ticker symbol', 'sym', 'cusip/symbol', 'symbol/cusip'],
    'desc': ['description', 'security name', 'security', 'security description', 'name', 'fund name', 'fund', 'holding', 'investment', 'account',
             'subaccount', 'position', 'asset', 'investment name'],
    'cls': ['asset class', 'asset type', 'type', 'category', 'class', 'security type', 'sector'],
    'qty': ['quantity', 'shares', 'units', 'qty', 'share quantity', 'shares/units', 'par'],
    'price': ['price', 'last price', 'unit price', 'market price', 'closing price', 'nav', 'share price'],
    'value': ['market value', 'value', 'balance', 'current value', 'amount', 'ending balance', 'total value', 'mkt value', 'ending value'],
    'cost': ['cost basis', 'cost', 'total cost', 'adjusted cost', 'book cost'],
    'pct': ['% of account', '% acct', '% of portfolio', 'allocation', 'weight', '% of total', 'percent', 'portfolio %', '%', '% of assets'],
    'gl': ['unrealized g/l', 'unrealized gain/loss', 'gain/loss', 'unrealized gain', 'g/l'],
    'date': ['date', 'trade date', 'settlement date', 'transaction date', 'activity date', 'process date'],
    'act': ['activity', 'transaction', 'action', 'transaction type', 'activity type'],
}
_SYN = {s: f for f, ss in FIELD_SYNONYMS.items() for s in ss}
NUMERIC_FIELDS = ('qty', 'price', 'value', 'cost', 'pct', 'gl')

MONTHS = {m: i + 1 for i, m in enumerate(['january', 'february', 'march', 'april', 'may', 'june', 'july', 'august', 'september', 'october', 'november', 'december'])}
MONTHS.update({k[:3]: v for k, v in list(MONTHS.items())})


@dataclass
class Seg:
    start: int
    end: int
    text: str
    conf: float = 0.97


@dataclass
class Line:
    text: str
    segs: list[Seg]
    conf: float = 0.97
    page: int = 1


@dataclass
class Statement:
    id: str
    file: str
    kind: str
    method: str
    institution: str | None = None
    owner: str | None = None
    account: str | None = None
    account_type: str | None = None
    as_of: str | None = None
    stated_total: float | None = None
    holdings: list[dict] = field(default_factory=list)
    allocation: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)
    page_map: list[dict] = field(default_factory=list)
    charts: list[dict] = field(default_factory=list)
    sections: list[dict] = field(default_factory=list)
    timings: dict = field(default_factory=dict)
    assets: dict = field(default_factory=dict)
    table_total: float | None = None

    def to_dict(self):
        ext = round(sum(h['value'] or 0 for h in self.holdings), 2)
        rec = None
        if self.stated_total:
            diff = round(ext - self.stated_total, 2)
            rec = {'stated': self.stated_total, 'extracted': ext, 'diff': diff, 'ok': abs(diff) <= max(1.0, self.stated_total * 0.0005)}
        confs = [h['confidence'] for h in self.holdings]
        return {'id': self.id, 'file': self.file, 'kind': self.kind, 'method': self.method, 'institution': self.institution, 'owner': self.owner,
                'account': self.account, 'account_type': self.account_type, 'account_category': account_category(self.account_type, self.lines),
                'as_of': self.as_of, 'stated_total': self.stated_total, 'extracted_total': ext, 'reconciliation': rec,
                'holdings': self.holdings, 'allocation': self.allocation, 'warnings': self.warnings,
                'confidence': round(min(confs), 2) if confs else 0.0, 'needs_review': sum(1 for h in self.holdings if h['review']),
                'pages': len(self.page_map) or 1, 'page_map': self.page_map, 'charts': self.charts, 'sections': self.sections,
                'table_total': self.table_total, 'timings': self.timings, 'text': '\n'.join(self.lines), '_assets': self.assets}


# ---------------------------------------------------------------------------------------------------------------- helpers
SEG_SPLIT = re.compile(r'\S+(?: \S+)*')          # runs separated by 2+ spaces
NUM_SEG = re.compile(r'^[\s$€£(+\-]*[\dOoSlI,.]*\d[\dOoSlI,.]*\)?\s*%?$')
TICKER = re.compile(r'^[A-Z][A-Z0-9]{0,4}(?:[.\-][A-Z0-9]{1,2})?$')
DASHES = {'-', '--', '—', '–', 'n/a', 'N/A', ''}


def segs_from_text(line: str, conf: float = 0.97) -> list[Seg]:
    return [Seg(m.start(), m.end(), m.group(0), conf) for m in re.finditer(r'\S+(?: \S+)*', line.replace('\t', '    '))]


def parse_number(s) -> float | None:
    """'$1,234.50' -> 1234.5; '(12.00)' -> -12; '43.9%' -> 43.9. OCR look-alikes inside numbers are repaired."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip()
    if t in DASHES:
        return None
    neg = t.startswith('(') and t.endswith(')') or t.startswith('-')
    t = t.strip('()$€£+-% ').replace('$', '')
    t = t.translate(str.maketrans({'O': '0', 'o': '0', 'l': '1', 'I': '1', 'S': '5'}))
    t = t.replace(',', '')
    if t.count('.') > 1:  # "1.234.56" OCR artefact: keep last dot
        head, _, tail = t.rpartition('.')
        t = head.replace('.', '') + '.' + tail
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def is_numeric(s: str) -> bool:
    return bool(NUM_SEG.match(s.strip())) and any(c.isdigit() for c in s)


def fix_ticker(s: str) -> tuple[str | None, bool]:
    """Return (ticker, corrected?) for a symbol cell. OCR often reads I as l or 1 and O as 0."""
    t = s.strip().strip('*')
    if t in DASHES or t.upper() in ('CASH', 'MMF', 'SWEEP'):
        return None, False
    if TICKER.match(t):
        return t, False
    if 1 <= len(t) <= 6 and re.match(r'^[A-Za-z0-9.\-]+$', t):
        u = t.replace('l', 'I').replace('|', 'I').upper()
        if re.search(r'[A-Z]', u):
            u2 = re.sub(r'(?<=[A-Z])0|0(?=[A-Z])', 'O', u)
            if TICKER.match(u2):
                return u2, True
    return None, False


def parse_date(s: str) -> str | None:
    s = s.strip()
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        return date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat()
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return m.group(0)
    m = re.search(r'([A-Za-z]{3,9})\.? (\d{1,2}),? (\d{4})', s)
    if m and m.group(1).lower() in MONTHS:
        return date(int(m.group(3)), MONTHS[m.group(1).lower()], int(m.group(2))).isoformat()
    return None


def header_field(text: str) -> str | None:
    t = re.sub(r'\s+', ' ', text.strip().lower().rstrip(':'))
    if t in _SYN:
        return _SYN[t]
    for s, f in sorted(_SYN.items(), key=lambda kv: -len(kv[0])):
        if len(s) > 3 and t.startswith(s):
            return f
    return None


# ---------------------------------------------------------------------------------------------------------------- readers
def read_pdf(path: Path, info: dict | None = None) -> tuple[list[Line], str]:
    """Text PDFs: the text layer, page by page. Scanned PDFs: two-pass OCR (see module notes)."""
    import time
    info = info if info is not None else {}
    t0 = time.time()
    txt = subprocess.run(['pdftotext', '-layout', str(path), '-'], capture_output=True, text=True, timeout=120).stdout
    pages_txt = txt.split('\f')
    if pages_txt and not pages_txt[-1].strip():
        pages_txt = pages_txt[:-1]
    if len(re.sub(r'\s', '', txt)) > 40 * max(1, len(pages_txt)) * 0.2:
        lines = []
        for page_no, page in enumerate(pages_txt, 1):
            for raw in page.splitlines():
                if raw.strip():
                    lines.append(Line(raw.rstrip(), segs_from_text(raw), 0.97, page_no))
        info['pages'] = len(pages_txt)
        info['visual'] = _pdf_visuals(path)
        info['timings'] = {'read': round(time.time() - t0, 2)}
        return lines, 'pdf-text'
    return _read_scanned(path, info), 'pdf-scan'


def _pdf_visuals(path: Path) -> dict:
    """Pages that carry drawings or images (possible charts)."""
    out = {}
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            for i, pg in enumerate(pdf.pages, 1):
                filled = [r for r in pg.rects if r.get('fill')]
                out[i] = {'images': len(pg.images), 'rects': len(filled)}
    except Exception:
        pass
    return out


def _n_pages(path: Path) -> int:
    r = subprocess.run(['pdfinfo', str(path)], capture_output=True, text=True, timeout=30).stdout
    m = re.search(r'Pages:\s+(\d+)', r)
    return int(m.group(1)) if m else 1


def _render(path: Path, page: int, dpi: int, color: bool = False):
    from PIL import Image
    with tempfile.TemporaryDirectory() as td:
        args = ['pdftoppm', '-r', str(dpi), '-f', str(page), '-l', str(page), '-png'] + ([] if color else ['-gray']) + [str(path), f'{td}/p']
        subprocess.run(args, check=True, timeout=120)
        im = Image.open(next(Path(td).glob('p*.png')))
        im.load()
        return im


def _colour_share(im) -> float:
    import numpy as np
    a = np.asarray(im.convert('RGB').resize((im.width // 4, im.height // 4))).astype(int)
    mx, mn = a.max(axis=2), a.min(axis=2)
    return float(((mx - mn) / np.maximum(mx, 1) > 0.25).mean())


def _read_scanned(path: Path, info: dict) -> list[Line]:
    import time
    from concurrent.futures import ThreadPoolExecutor
    import os
    t0 = time.time()
    n = _n_pages(path)
    workers = max(1, min(4, os.cpu_count() or 1))

    def quick(pg):
        im = _render(path, pg, 80, color=True)
        return pg, ocr_image(im.convert('L'), page=pg, fast=True), _colour_share(im)
    with ThreadPoolExecutor(workers) as ex:
        first = sorted(ex.map(quick, range(1, n + 1)))
    t1 = time.time()
    by_page = {pg: ls for pg, ls, _ in first}
    visual = {pg: {'images': 1 if share > 0.01 else 0, 'rects': 0, 'colour': round(share, 4)} for pg, _, share in first}
    types = classify_pages(by_page, visual, n)
    # full resolution for summary and holdings pages, and for any unclassed page that looks like a table
    need = [pg for pg in range(1, n + 1) if types[pg]['type'] in ('summary', 'holdings') or pg == 1
            or (types[pg]['type'] == 'other' and types[pg]['numeric_rows'] >= 3)]

    def full(pg):
        return pg, ocr_image(_render(path, pg, 250), page=pg)
    with ThreadPoolExecutor(workers) as ex:
        for pg, ls in ex.map(full, need):
            by_page[pg] = ls
    info['pages'] = n
    info['visual'] = visual
    info['types'] = types
    info['full_ocr_pages'] = need
    info['timings'] = {'triage_ocr': round(t1 - t0, 1), 'full_ocr': round(time.time() - t1, 1), 'pages_full_ocr': len(need), 'pages_total': n}
    return [l for pg in range(1, n + 1) for l in by_page[pg]]


def ocr_image(path, page: int = 1, fast: bool = False) -> list[Line]:
    import pytesseract
    from PIL import Image, ImageOps
    img = path if isinstance(path, Image.Image) else Image.open(path)
    img = ImageOps.exif_transpose(img).convert('L')
    if not fast:
        if img.width < 1600:
            img = img.resize((int(img.width * 1600 / img.width), int(img.height * 1600 / img.width)), Image.LANCZOS)
        img = ImageOps.autocontrast(img, cutoff=1)
        img = deskew(img)
    d = pytesseract.image_to_data(img, config='--psm 6 -c preserve_interword_spaces=1', output_type=pytesseract.Output.DICT)
    words = [dict(t=d['text'][i], x=d['left'][i], w=d['width'][i], y=d['top'][i], h=d['height'][i], c=float(d['conf'][i]),
                  k=(d['block_num'][i], d['par_num'][i], d['line_num'][i])) for i in range(len(d['text'])) if d['text'][i].strip()]
    if not words:
        return []
    charw = sorted(w['w'] / max(1, len(w['t'])) for w in words)[len(words) // 2]
    groups: dict = {}
    for w in words:
        groups.setdefault(w['k'], []).append(w)
    lines = []
    for k in sorted(groups, key=lambda k: min(w['y'] for w in groups[k])):
        ws = sorted(groups[k], key=lambda w: w['x'])
        # place words on a character grid so columns line up across lines, like pdftotext -layout
        text, segs, cur = '', [], None
        for w in ws:
            col = int(round(w['x'] / charw))
            gap_px = w['x'] - (cur['x'] + cur['w']) if cur else None
            new_seg = cur is None or gap_px > 1.6 * charw
            if cur is not None:
                col = max(col, len(text) + (3 if new_seg else 1))
                if not new_seg:
                    col = len(text) + 1
            text = text.ljust(col) + w['t']
            if new_seg:
                segs.append(Seg(col, len(text), w['t'], [w['c']]))
            else:
                segs[-1].end = len(text); segs[-1].text += ' ' + w['t']; segs[-1].conf.append(w['c'])
            cur = w
        for s in segs:
            s.conf = round(max(0.0, min(1.0, sum(s.conf) / len(s.conf) / 100)), 3)
        lines.append(Line(text, segs, round(min(s.conf for s in segs), 3), page))
    return lines


def deskew(img):
    """Small-angle deskew: pick the rotation that makes the row profile sharpest."""
    import numpy as np
    from PIL import Image
    small = img.resize((img.width // 4, img.height // 4))
    a = 255 - np.asarray(small, dtype=np.float32)
    best, best_score = 0.0, -1
    for ang in [x / 2 for x in range(-8, 9)]:
        r = Image.fromarray(a.astype('uint8')).rotate(ang, resample=Image.BILINEAR)
        prof = np.asarray(r, dtype=np.float32).sum(axis=1)
        score = float(((prof[1:] - prof[:-1]) ** 2).sum())
        if score > best_score:
            best, best_score = ang, score
    return img.rotate(best, resample=Image.BICUBIC, fillcolor=255, expand=True) if best else img


def read_cells(path: Path, kind: str) -> list[list[str]]:
    if kind == 'csv':
        raw = path.read_bytes().decode('utf-8-sig', errors='replace')
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(raw))]
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets[:3]:
        for r in ws.iter_rows(values_only=True):
            rows.append(['' if v is None else (v if isinstance(v, (int, float)) else str(v).strip()) for v in r])
        rows.append([])
    return rows


# ---------------------------------------------------------------------------------------------------------------- parsing
META_PATTERNS = {
    'owner': re.compile(r'(?:prepared for|account holder|contract owner|account owner|client|owner|participant|investor)\s*[:,]\s*(.+)', re.I),
    'account_type': re.compile(r'account type\s*[:,]\s*(.+)', re.I),
    'account': re.compile(r'(?:x{2,}|\*{2,}|•{2,})[\s-]*(\d{3,5})\b', re.I),
}
TOTAL_LABEL = re.compile(r'(total account value|contract value|account value|plan balance|total portfolio value|total value|total deposits|account total|'
                         r'total market value|portfolio value|total balance|ending value)', re.I)
DATE_CONTEXT = re.compile(r'(as of|statement period|period ending|statement date|balance as of|period|date)', re.I)
TOTAL_ROW = re.compile(r'^(?:grand |sub)?total(?: (?:account|portfolio|holdings|deposits|value|market value|assets))?:?$|^(?:account|portfolio) total:?$', re.I)
ALLOC_LABEL = re.compile(r'^(?:asset )?allocation\b|\ballocation$|^asset mix|^portfolio mix', re.I)


def money_in(s: str) -> float | None:
    m = re.search(r'\$\s?([\d,]+(?:\.\d{2})?)', s)
    return parse_number(m.group(1)) if m else None


def parse_meta(st: Statement, lines: list[str]):
    head = [l for l in lines[:40]]
    for l in head:
        flat = re.sub(r'\s{2,}', '  ', l.strip())
        if not st.owner:
            m = META_PATTERNS['owner'].search(flat)
            if m:
                st.owner = re.split(r'\s{2,}|\s\|\s', m.group(1).strip())[0].strip(' ,')
        if not st.account_type:
            m = META_PATTERNS['account_type'].search(flat)
            if m:
                st.account_type = re.split(r'\s{2,}', m.group(1).strip())[0]
        if not st.account:
            m = META_PATTERNS['account'].search(flat)
            if m:
                st.account = m.group(1)[-4:]
        if not st.as_of and DATE_CONTEXT.search(flat):
            ds = re.findall(r'\d{1,2}/\d{1,2}/\d{4}|\d{4}-\d{2}-\d{2}|[A-Za-z]{3,9}\.? \d{1,2},? \d{4}', flat)
            if ds:
                st.as_of = parse_date(ds[-1])
        if st.stated_total is None and TOTAL_LABEL.search(flat):
            tail = flat[TOTAL_LABEL.search(flat).end():]
            cands = [parse_number(t) for t in re.findall(r'\$?[\d,]+(?:\.\d+)?%?', tail) if not t.endswith('%')]
            v = money_in(tail) or max([c for c in cands if c is not None] or [0])
            if v and v > 0:
                st.stated_total = v
        if not st.owner and ' | log out' in flat.lower():
            st.owner = flat.lower().split(' | log out')[0].split('  ')[-1].strip().title()
        if not st.owner:
            m = re.search(r'overview\s*[-–]\s*([A-Z][A-Za-z]+ (?:Household|Family(?: Trust)?))', flat)
            if m:
                st.owner = m.group(1)
    # institution: first meaningful line, first segment, stripped of " - Holdings export" style suffixes
    for l in head:
        seg = re.split(r'\s{2,}', l.strip())[0]
        if len(seg) > 3 and not re.match(r'^(my account|home|log ?in|statement|page \d)', seg, re.I) and not is_numeric(seg):
            st.institution = re.split(r'\s[-–]\s', seg)[0].strip()
            break
    # a lone large dollar figure just after a "balance" caption (portal screenshots)
    if st.stated_total is None:
        for i, l in enumerate(head):
            if re.fullmatch(r'\s*\$[\d,]+\.\d{2}\s*', l) and i > 0 and re.search(r'balance|value|total|portfolio', ' '.join(head[max(0, i - 2):i + 1]), re.I):
                st.stated_total = money_in(l)
                break
        if st.stated_total is None:
            for l in head[:12]:
                if re.fullmatch(r'\s*\$[\d,]+\.\d{2}\s*', l):
                    st.stated_total = money_in(l)
                    break
    if not st.as_of:
        for l in lines:
            ds = re.findall(r'\d{1,2}/\d{1,2}/\d{4}|[A-Za-z]{3,9}\.? \d{1,2},? \d{4}', l)
            if ds and DATE_CONTEXT.search(l):
                st.as_of = parse_date(ds[-1]); break
    if not st.account_type:
        st.account_type = guess_account_type('\n'.join(lines))


def guess_account_type(text: str) -> str | None:
    t = text.lower()
    for pat, name in [(r'401\(k\)|403\(b\)|457\b|retirement plan', '401(k) plan'), (r'roth ira', 'Roth IRA'), (r'\bira\b|individual retirement', 'IRA'),
                      (r'annuity', 'Variable annuity'), (r'joint', 'Joint taxable'), (r'deposit|savings|checking|certificate of deposit', 'Deposit accounts'),
                      (r'trust', 'Trust account'), (r'brokerage|individual', 'Individual brokerage (taxable)')]:
        if re.search(pat, t):
            return name
    return 'Investment account'


def account_category(account_type: str | None, lines: list[str]) -> str:
    t = (account_type or '').lower() + ' ' + ' '.join(lines[:15]).lower()
    if re.search(r'401\(k\)|403\(b\)|457|retirement plan', t):
        return 'employer-plan'
    if 'annuity' in t:
        return 'annuity'
    if re.search(r'\bira\b|individual retirement', t):
        return 'ira'
    if re.search(r'deposit|bank|savings', t) and not re.search(r'brokerage|securities', t):
        return 'bank'
    return 'taxable'


def map_header(segs: list[Seg]) -> dict | None:
    cols = [(header_field(s.text), s) for s in segs]
    known = [(f, s) for f, s in cols if f]
    fields = {f for f, _ in known}
    if len(known) >= 2 and (fields & {'value', 'qty', 'pct', 'price'}) and (fields & {'desc', 'sym'}):
        return {'cols': [(f or '?', s) for f, s in cols], 'activity': bool(fields & {'date', 'act'})}
    return None


def assign_cells(segs: list[Seg], header: dict) -> dict:
    """Map a data line's segments onto header columns.

    Type first (a percent goes to the % column, a ticker-looking cell to the symbol column, text to the description),
    then position (largest overlap with the header cell, else the nearest header cell of the right type)."""
    cols = header['cols']
    numeric_cols = [(f, s) for f, s in cols if f in NUMERIC_FIELDS]
    out: dict = {}
    nums = []
    for sg in segs:
        txt = sg.text.strip()
        if is_numeric(txt) or txt in DASHES:
            if txt.endswith('%') and any(f == 'pct' for f, _ in cols):
                out['pct'] = (txt, sg.conf); continue
            if txt in DASHES:
                # a dash in the symbol column means "no symbol"
                near = min(cols, key=lambda c: abs((c[1].start + c[1].end) / 2 - (sg.start + sg.end) / 2))
                if near[0] == 'sym':
                    out['sym'] = ('', sg.conf)
                continue
            nums.append(sg)
            continue
        tk, _ = fix_ticker(txt)
        sym_col = next((s for f, s in cols if f == 'sym'), None)
        if sym_col is not None and tk and 'sym' not in out and (len(txt) <= 6) and (txt.upper() == txt or _overlap(sg, sym_col) > 0):
            out['sym'] = (txt, sg.conf); continue
        f = _best_col(sg, cols, allowed={'desc', 'cls', 'sym'})
        if f == 'sym' and not tk:
            f = 'desc'
        if f in out:
            prev = out[f]
            out[f] = (prev[0] + ' ' + txt, min(prev[1], sg.conf))
        else:
            out[f or 'desc'] = (txt, sg.conf)
    taken = set(out)
    num_fields = [f for f, _ in numeric_cols if f != 'pct' or 'pct' not in taken]
    if len(nums) == len([f for f in num_fields if f not in taken]):
        for sg, f in zip(nums, [f for f in num_fields if f not in taken]):
            out[f] = (sg.text, sg.conf)
    else:
        for sg in nums:
            f = _best_col(sg, [(f, s) for f, s in numeric_cols if f not in out], allowed=set(NUMERIC_FIELDS))
            if f:
                out[f] = (sg.text, sg.conf)
    return out


def _overlap(a: Seg, b: Seg) -> int:
    return max(0, min(a.end, b.end) - max(a.start, b.start))


def _best_col(sg: Seg, cols, allowed) -> str | None:
    cand = [(f, s) for f, s in cols if f in allowed]
    if not cand:
        return None
    best = max(cand, key=lambda c: (_overlap(sg, c[1]), -abs((c[1].start + c[1].end) / 2 - (sg.start + sg.end) / 2)))
    return best[0]


def build_holding(row: dict, idx: int, base_conf: float, source: str) -> dict | None:
    flags = []
    sym_raw = row.get('sym', ('', 1))[0]
    sym, corrected = fix_ticker(str(sym_raw)) if sym_raw not in (None, '') else (None, False)
    if corrected:
        flags.append(f'Symbol read as "{sym_raw}", corrected to {sym}')
    desc = str(row.get('desc', ('', 1))[0] or '').strip()
    if not sym and not desc:
        return None
    vals = {f: parse_number(row[f][0]) if f in row else None for f in NUMERIC_FIELDS}
    confs = [c for f, (v, c) in row.items() if v not in ('', None)]
    conf = min([base_conf] + [c for c in confs if isinstance(c, float)])
    if vals['value'] is None and vals['qty'] is not None and vals['price'] is not None:
        vals['value'] = round(vals['qty'] * vals['price'], 2)
        flags.append('Market value not shown; computed as quantity x price')
        conf -= 0.05
    if vals['value'] is None:
        return None
    if vals['qty'] and vals['price']:
        calc = vals['qty'] * vals['price']
        # bonds are often priced per 100 of par
        if abs(calc - vals['value']) > max(1.0, 0.005 * vals['value']) and abs(calc / 100 - vals['value']) > max(1.0, 0.005 * vals['value']):
            flags.append(f'Quantity x price = {calc:,.2f} does not match market value {vals["value"]:,.2f}')
            conf -= 0.3
    if corrected:
        conf -= 0.1
    conf = round(max(0.0, min(1.0, conf)), 2)
    if desc.lower() in ('cash', 'cash & sweep') and not sym:
        desc = desc or 'Cash'
    return {'row': idx, 'sym': sym, 'desc': desc or sym, 'cls_hint': (row.get('cls') or (None,))[0], 'qty': vals['qty'], 'price': vals['price'],
            'value': vals['value'], 'cost': vals['cost'], 'pct': vals['pct'], 'confidence': conf, 'review': conf < 0.85 or bool(flags),
            'flags': flags, 'source': source}


DISC_KW = re.compile(r'important information|disclosure|glossary|privacy|terms and conditions|about your (statement|account)|product information|'
                     r'definitions|notice|protection program|complaints', re.I)
ACT_KW = re.compile(r'\b(activity|transactions?|trade date|settlement date|purchases and sales)\b', re.I)
HOLD_KW = re.compile(r'\b(holdings|positions|portfolio detail|investment detail|your investments|securities held)\b', re.I)
CHART_KW = re.compile(r'allocation|by sector|asset mix|diversification|by asset class|composition|breakdown', re.I)
DATE_START = re.compile(r'^\s*\d{1,2}/\d{1,2}(/\d{2,4})?\b')
SECTION_WORDS = re.compile(r'^(equities|equity|stocks|common stocks?|fixed income|bonds?|cash( and (cash )?equivalents)?|cash & equivalents|mutual funds|'
                           r'exchange[- ]traded( funds| and mutual funds)?|etfs?|alternative investments|alternatives|preferred|options|other)\b', re.I)
TABLE_END = re.compile(r'^(total (holdings|portfolio|account|assets|market value|investments|securities)|grand total|account total|portfolio total|total)\b:?\s*$', re.I)


def classify_pages(by_page: dict, visual: dict, n: int) -> dict:
    """Class every page from its text and drawings. Holdings continue onto pages that look like more of the same table."""
    out, prev = {}, None
    for pg in range(1, n + 1):
        ls = by_page.get(pg, [])
        texts = [l.text.strip() for l in ls if l.text.strip()]
        head = ' '.join(texts[:5])
        numeric_rows = sum(1 for l in ls if sum(is_numeric(s.text) for s in l.segs) >= 2)
        prose = sum(1 for t in texts if len(t.split()) >= 12 and sum(c.isdigit() for c in t) < len(t) * 0.12)
        dated = sum(1 for t in texts if DATE_START.match(t))
        hdrs = [h for h in (map_header(l.segs) for l in ls) if h]
        v = visual.get(pg, {})
        has_visual = v.get('images', 0) > 0 or v.get('rects', 0) >= 3
        title = ' '.join(texts[:3])
        if any(h['activity'] for h in hdrs) or (dated >= 4 and ACT_KW.search(head)) or (ACT_KW.search(title) and not HOLD_KW.search(title)):
            t, why = 'activity', 'transaction table' if hdrs or dated >= 4 else 'activity title'
        elif any(not h['activity'] for h in hdrs):
            t, why = 'holdings', 'holdings table header'
        elif HOLD_KW.search(title) and not DISC_KW.search(title):
            t, why = 'holdings', 'holdings title'
        elif prev == 'holdings' and numeric_rows >= 2 and prose < numeric_rows and not DISC_KW.search(head):
            t, why = 'holdings', 'continues the holdings table'
        elif has_visual and (CHART_KW.search(head) or numeric_rows < 4) and not DISC_KW.search(head):
            t, why = 'chart', 'graphics with an allocation title' if CHART_KW.search(head) else 'graphics'
        elif DISC_KW.search(head) or (prose >= 4 and numeric_rows < 3):
            t, why = 'disclosure', 'disclosure or prose page'
        elif pg == 1 or re.search(r'total (account|portfolio) value|account value|summary|overview', head, re.I):
            t, why = 'summary', 'account summary'
        elif HOLD_KW.search(head) and numeric_rows >= 3:
            t, why = 'holdings', 'holdings title with rows'
        else:
            t, why = 'other', 'not needed for holdings'
        out[pg] = {'page': pg, 'type': t, 'why': why, 'numeric_rows': numeric_rows}
        prev = t if t != 'chart' else prev
    return out


def parse_lines(st: Statement, lines: list[Line], types: dict | None = None):
    st.lines = [l.text for l in lines]
    st.page_titles = {}
    for l in lines:
        t = l.text.strip()
        if CHART_KW.search(t) and len(t.split()) <= 8 and not re.search(r'x{2,}|account|page \d', t, re.I):
            if l.page not in st.page_titles or len(t) < len(st.page_titles[l.page]):
                st.page_titles[l.page] = t
    pages = sorted({l.page for l in lines}) or [1]
    by_page: dict = {}
    for l in lines:
        by_page.setdefault(l.page, []).append(l)
    if types is None:
        types = classify_pages(by_page, {}, max(pages)) if len(pages) > 1 else {pages[0]: {'page': pages[0], 'type': 'summary', 'why': 'single page', 'numeric_rows': 0}}
    use = {pg for pg, t in types.items() if t['type'] in ('summary', 'holdings', 'chart')} or set(pages)
    meta_lines = [l.text for l in lines if l.page in use or l.page == 1]
    parse_meta(st, meta_lines)
    header = None
    in_alloc = False
    idx = 0
    section, sec_start = None, 0
    last_holding = None
    rows_on_page: dict = {}
    for pg in pages:
        ptype = types.get(pg, {}).get('type', 'summary')
        if ptype not in ('summary', 'holdings', 'chart'):
            header, in_alloc, last_holding = (header if ptype == 'chart' else None), False, None
            continue
        if ptype != 'holdings':
            header = None
        for ln in by_page.get(pg, []):
            flat = ln.text.strip()
            low = flat.lower()
            h = map_header(ln.segs)
            if h:
                header, in_alloc = (None if h['activity'] else h), False
                last_holding = None
                continue
            if not ln.segs:
                continue
            has_num = any(is_numeric(s.text) for s in ln.segs)
            if ALLOC_LABEL.search(flat) and not has_num and '%' not in flat:
                in_alloc, header = True, None
                continue
            # inline allocation: "Stocks 68%   Bonds 24%" or "Allocation: Cash 100%"
            inline = re.findall(r'([A-Za-z][A-Za-z &/\-]{1,40}?)\s+(\d{1,3}(?:\.\d+)?)%', re.sub(r'^.*?allocation\s*:\s*', '', flat, flags=re.I))
            if (in_alloc or re.match(r'^(asset )?allocation\s*:', low)) and inline and all(re.fullmatch(r'[A-Za-z][A-Za-z &/\-]{1,40}?\s+\d{1,3}(?:\.\d+)?%', s.text.strip())
                                                                                            or re.match(r'^allocation\s*:', s.text.strip(), re.I) for s in ln.segs):
                for lab, p in inline:
                    st.allocation.append({'label': lab.strip(), 'pct': float(p), 'value': None, 'source': 'text'})
                continue
            if in_alloc and has_num:
                txt = [s for s in ln.segs if not is_numeric(s.text)]
                pct = next((parse_number(s.text) for s in ln.segs if s.text.strip().endswith('%')), None)
                val = next((parse_number(s.text) for s in ln.segs if is_numeric(s.text) and not s.text.strip().endswith('%')), None)
                if txt and pct is not None:
                    st.allocation.append({'label': txt[0].text.strip(), 'pct': pct, 'value': val, 'source': 'text'})
                    continue
            if header is None:
                continue
            first_txt = next((s.text.strip() for s in ln.segs if not is_numeric(s.text)), '')
            nums = [parse_number(s.text) for s in ln.segs if is_numeric(s.text) and not s.text.strip().endswith('%')]
            # totals: a grand total ends the table; a section subtotal is checked and the table goes on
            rest = re.sub(r'^(sub)?total\s*', '', first_txt, flags=re.I)
            has_ticker = any(fix_ticker(sg.text.strip())[0] and sg.text.strip().upper() == sg.text.strip() and len(sg.text.strip()) <= 6 for sg in ln.segs[1:] if not is_numeric(sg.text))
            is_total = TOTAL_ROW.match(first_txt) or TABLE_END.match(first_txt) or (re.match(r'^(sub)?total\b', first_txt, re.I) and not has_ticker and (
                not rest or SECTION_WORDS.match(rest) or (section and rest.lower() == section.lower())))
            if is_total:
                if nums:
                    v = nums[0]
                    rows = st.holdings[sec_start:]
                    if TABLE_END.match(first_txt):
                        st.table_total = v
                        if st.stated_total is None:
                            st.stated_total = v
                        header = None
                    else:
                        got = round(sum(r['value'] for r in rows), 2)
                        st.sections.append({'label': first_txt, 'stated': v, 'rows': len(rows), 'sum': got, 'ok': abs(got - v) <= max(1.0, v * 0.0005), 'page': pg})
                        sec_start = len(st.holdings)
                last_holding = None
                continue
            if not has_num:
                # section heading, or the second line of a wrapped description
                if SECTION_WORDS.match(first_txt) and len(ln.segs) == 1 and len(first_txt) < 45:
                    if first_txt.isupper() or section is None:
                        section = first_txt.title() if first_txt.isupper() else first_txt
                    last_holding = None
                    continue
                if last_holding is not None and len(ln.segs) <= 2:
                    desc_col = next((s for f, s in header['cols'] if f == 'desc'), None)
                    indented = desc_col is not None and ln.segs[0].start >= desc_col.start + 2
                    if indented or re.search(r'\b(CUSIP|ISIN|Moody|S&P|Fitch|Callable|Maturity|Coupon|Rating)\b', flat, re.I):
                        last_holding['detail'] = (last_holding.get('detail', '') + ' ' + flat).strip()
                        m = re.search(r'\b(?:CUSIP|ISIN)[:\s#]*([A-Z0-9]{9,12})\b', flat)
                        if m:
                            last_holding['cusip'] = m.group(1)
                        continue
                continue
            cells = assign_cells(ln.segs, header)
            idx += 1
            base = min([s.conf for s in ln.segs if s.text.strip() not in DASHES] or [ln.conf])
            hld = build_holding(cells, idx, base, flat)
            if hld:
                hld['page'] = pg
                if section and not hld.get('cls_hint'):
                    hld['cls_hint'] = section
                hld['section'] = section
                st.holdings.append(hld)
                rows_on_page[pg] = rows_on_page.get(pg, 0) + 1
                last_holding = hld
    for pg in sorted(types):
        t = dict(types[pg])
        t['rows'] = rows_on_page.get(pg, 0)
        st.page_map.append(t)
    for sct in st.sections:
        if not sct['ok']:
            st.warnings.append(f"{sct['label']}: rows add up to ${sct['sum']:,.2f}, statement shows ${sct['stated']:,.2f}")


def parse_cells(st: Statement, rows: list[list]):
    st.lines = [' , '.join(str(c) for c in r if c not in ('', None)) for r in rows]
    hi, cols = None, None
    for i, r in enumerate(rows[:40]):
        fields = [header_field(str(c)) if c not in ('', None) else None for c in r]
        if sum(1 for f in fields if f) >= 3 and set(fields) & {'value', 'qty'}:
            hi, cols = i, fields
            break
    meta_lines = []
    for r in rows[:hi if hi is not None else 10]:
        cells = [str(c) for c in r if c not in ('', None)]
        if len(cells) == 2:
            meta_lines.append(f'{cells[0]}: {cells[1]}')
        elif cells:
            meta_lines.append('  '.join(cells))
    parse_meta(st, meta_lines + st.lines[(hi or 0):])
    if not st.account_type or st.account_type in ('Investment account',):
        for l in meta_lines:
            m = re.match(r'account\s*:\s*(.+?)\s+(?:x{2,}|\*{2,})', l, re.I)
            if m:
                st.account_type = m.group(1).strip()
    if hi is None:
        st.warnings.append('No holdings table header found')
        return
    idx = 0
    for r in rows[hi + 1:]:
        if not any(c not in ('', None) for c in r):
            break
        row = {}
        for f, c in zip(cols, r):
            if f and c not in ('', None):
                row[f] = (c, 0.99)
        d = str(row.get('desc', ('',))[0]).lower()
        if TOTAL_ROW.match(d.strip()) or TOTAL_ROW.match(str(r[0]).strip()):
            if st.stated_total is None and 'value' in row:
                st.stated_total = parse_number(row['value'][0])
            break
        idx += 1
        h = build_holding(row, idx, 0.99, ' | '.join(str(c) for c in r if c not in ('', None)))
        if h:
            st.holdings.append(h)


# ---------------------------------------------------------------------------------------------------------------- entry
def extract(path: str | Path, sid: str | None = None, display_name: str | None = None) -> dict:
    path = Path(path)
    name = display_name or path.name
    kind = SUPPORTED.get(path.suffix.lower())
    st = Statement(id=sid or path.stem, file=name, kind=kind or 'unknown', method='')
    if not kind:
        st.warnings.append(f'Unsupported file type {path.suffix}')
        return st.to_dict()
    import time
    if kind == 'pdf':
        info: dict = {}
        lines, st.kind = read_pdf(path, info)
        n = info.get('pages', 1)
        if st.kind == 'pdf-text':
            st.method = 'PDF text layer'
            by_page: dict = {}
            for l in lines:
                by_page.setdefault(l.page, []).append(l)
            types = classify_pages(by_page, info.get('visual', {}), n)
        else:
            types = info['types']
            st.method = f"OCR of scanned pages ({len(info['full_ocr_pages'])} of {n} at full resolution)" if n > 1 else 'OCR of scanned pages'
        st.timings.update(info.get('timings', {}))
        t0 = time.time()
        parse_lines(st, lines, types)
        st.timings['parse'] = round(time.time() - t0, 2)
        t0 = time.time()
        read_charts(st, path, types, info)
        st.timings['charts'] = round(time.time() - t0, 2)
    elif kind == 'image':
        st.method = 'OCR of image'
        parse_lines(st, ocr_image(path))
        if not st.allocation:
            from PIL import Image
            read_image_chart(st, Image.open(path))
    elif kind == 'text':
        st.method = 'Plain text'
        parse_lines(st, [Line(l, segs_from_text(l)) for l in path.read_text(errors='replace').splitlines() if l.strip()])
    else:
        st.method = 'Spreadsheet cells' if kind == 'xlsx' else 'CSV cells'
        parse_cells(st, read_cells(path, kind))
    d = st.to_dict()
    if len(d['page_map']) > 1:
        skipped = [p['page'] for p in d['page_map'] if p['type'] in ('disclosure', 'activity', 'other')]
        if skipped:
            d['notes'] = [f"{len(skipped)} of {len(d['page_map'])} pages not needed for holdings (activity, fee tables, disclosures): pages {_ranges(skipped)}"]
    rec = d['reconciliation']
    if rec and not rec['ok']:
        d['warnings'].append(f"Rows add up to ${rec['extracted']:,.2f}, statement total is ${rec['stated']:,.2f} (difference ${rec['diff']:,.2f})")
    if not d['holdings']:
        d['warnings'].append('No holdings found: enter them by hand or upload a clearer copy')
    if d['allocation']:
        tot = sum(a['pct'] for a in d['allocation'])
        if abs(tot - 100) > 1.5:
            d['warnings'].append(f'Allocation percentages add up to {tot:.1f}%')
    return d


def _ranges(nums):
    out, start, prev = [], None, None
    for x in sorted(nums):
        if start is None:
            start = prev = x
        elif x == prev + 1:
            prev = x
        else:
            out.append(f'{start}' if start == prev else f'{start}-{prev}'); start = prev = x
    if start is not None:
        out.append(f'{start}' if start == prev else f'{start}-{prev}')
    return ', '.join(out)


def read_charts(st: Statement, path: Path, types: dict, info: dict):
    """Measure allocation charts on chart and summary pages."""
    from . import charts as C
    cand = [pg for pg, t in types.items() if t['type'] == 'chart' or (t['type'] == 'summary' and info.get('visual', {}).get(pg, {}).get('images'))]
    found = []
    if st.kind == 'pdf-text' and cand:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            for pg in cand:
                page = pdf.pages[pg - 1]
                title = _page_title(st, pg)
                bars = C.vector_bars(page, pg, title)
                found += bars
                if page.images:
                    full = _render(path, pg, 150, color=True)
                    sx, sy = full.width / float(page.width), full.height / float(page.height)
                    for im in page.images:
                        crop = full.crop((int(im['x0'] * sx), int(im['top'] * sy), int(im['x1'] * sx), int(im['bottom'] * sy)))
                        r = C.raster_pie(crop, pg, title)
                        if r:
                            found.append(r)
    elif st.kind == 'pdf-scan':
        for pg in cand:
            r = C.raster_pie(_render(path, pg, 150, color=True), pg, _page_title(st, pg))
            if r:
                found.append(r)
    _attach_charts(st, found)


def read_image_chart(st: Statement, img):
    from . import charts as C
    r = C.raster_pie(img.convert('RGB'), 1, None)
    _attach_charts(st, [r] if r else [])


def _attach_charts(st: Statement, found):
    from . import llm
    for k, r in enumerate(found):
        key = f'chart-{r.page}-{k}'
        if r.image:
            st.assets[key] = r.image
        d = r.to_dict()
        d['asset'] = key if r.image else None
        if llm.enabled() and r.confidence < 0.8 and r.image:
            try:
                d = llm.check_chart(d, r.image)
            except Exception as e:
                d['note'] += f'; language model unavailable ({e.__class__.__name__})'
        st.charts.append(d)
    # an allocation shown only as a chart becomes the statement's allocation
    if not st.allocation:
        pie = next((c for c in st.charts if c['type'] == 'pie'), None)
        if pie:
            st.allocation = [{'label': v['label'], 'pct': v['pct'], 'value': None, 'source': f"chart, page {pie['page']}"} for v in pie['values']]


def _page_title(st: Statement, pg: int) -> str | None:
    return getattr(st, 'page_titles', {}).get(pg)


def tools_available() -> dict:
    ok = {'pdftotext': bool(shutil.which('pdftotext')), 'pdftoppm': bool(shutil.which('pdftoppm')), 'tesseract': bool(shutil.which('tesseract'))}
    try:
        import openpyxl  # noqa
        ok['openpyxl'] = True
    except ImportError:
        ok['openpyxl'] = False
    return ok
